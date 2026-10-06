import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    """Sets background color for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def create_milestones_doc(output_path):
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Color Palette
    PRIMARY_COLOR = RGBColor(15, 34, 64)       # Deep Navy
    SECONDARY_COLOR = RGBColor(30, 90, 150)   # Steel Blue
    DARK_TEXT = RGBColor(40, 40, 40)          # Charcoal Text
    LIGHT_BG = "F4F6F9"                       # Light Grey Blue

    # Base Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = DARK_TEXT

    # --- Document Header Title ---
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("HTC Core Enterprise Platform")
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    p_subtitle = doc.add_paragraph()
    run_sub = p_subtitle.add_run("Development to Deployment Milestones Plan")
    run_sub.font.size = Pt(16)
    run_sub.font.color.rgb = SECONDARY_COLOR
    run_sub.font.bold = True

    # Metadata Block
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(14)
    p_meta.add_run("Prepared For: ").bold = True
    p_meta.add_run("Project Stakeholders & Capstone Evaluation Panel\n")
    p_meta.add_run("Document Version: ").bold = True
    p_meta.add_run("1.0  |  ")
    p_meta.add_run("Status: ").bold = True
    p_meta.add_run("Approved for Sign-Off & Deployment")

    # Divider
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(12)
    run_div = p_div.add_run("―" * 48)
    run_div.font.color.rgb = RGBColor(200, 200, 200)

    # --- Section: Executive Summary ---
    h1 = doc.add_heading(level=1)
    run_h1 = h1.add_run("1. Executive Summary")
    run_h1.font.color.rgb = PRIMARY_COLOR
    run_h1.font.size = Pt(14)
    run_h1.font.bold = True

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(14)
    p_exec.add_run(
        "This document outlines the formal development, testing, and cloud deployment milestones for the "
        "HTC Core Enterprise System. Designed to ensure complete alignment between academic evaluation standards "
        "and enterprise software best practices, this roadmap breaks the project into six distinct, measurable "
        "milestones. Each milestone features explicit deliverables, acceptance criteria, technical validation standards, "
        "and stakeholder approval metrics."
    )

    # --- Section: Milestone Roadmap Breakdown ---
    h2 = doc.add_heading(level=1)
    run_h2 = h2.add_run("2. Detailed Milestone Specifications")
    run_h2.font.color.rgb = PRIMARY_COLOR
    run_h2.font.size = Pt(14)
    run_h2.font.bold = True

    milestones_data = [
        {
            "num": "Milestone 1",
            "title": "System Architecture, Database Schema & Core Security Framework",
            "focus": "Establishing foundational infrastructure, multi-role access controls, and database design.",
            "objective": "Build a resilient, scalable backend foundation with strict Role-Based Access Control (RBAC) and immutable data models.",
            "deliverables": [
                "Unified Django application architecture with modular app structure (operations, config, templates).",
                "Relational Database Schema (PostgreSQL/SQLite) with normalized models for Users, Roles, Audit Logs, and Workflows.",
                "Granular Permission Engine supporting Admin, Procurement Specialist, Logistics Manager, and Finance Officer roles.",
                "Environment variable management (.env) for secrets, API credentials, and database settings."
            ],
            "criteria": [
                "100% of user roles properly enforce route guard restrictions and permission checks.",
                "Database migrations execute cleanly without circular dependencies or orphaned foreign keys.",
                "All sensitive credentials are isolated from source control."
            ]
        },
        {
            "num": "Milestone 2",
            "title": "Procurement, Vendor & Inventory/Logistics Module Implementation",
            "focus": "Automating core supply chain workflows, purchase order lifecycle, and stock tracking.",
            "objective": "Develop end-to-end management tools for requisitions, vendor proposals, purchase orders, inventory levels, and dispute tracking.",
            "deliverables": [
                "Vendor Directory & Quotation Management Module.",
                "Automated Requisition-to-Purchase-Order approval pipeline.",
                "Multi-warehouse Inventory Tracker with real-time stock adjustments.",
                "Logistics & Returns Management system for logging damaged or disputed shipments."
            ],
            "criteria": [
                "System successfully processes end-to-end Requisition -> PO -> Stock Receiving workflow.",
                "Inventory updates automatically upon stock movement or dispute approval.",
                "Multi-file attachment and CSV/Excel import validators function without memory leakage."
            ]
        },
        {
            "num": "Milestone 3",
            "title": "Financial Accounting, Loan Engine & Audit Logging",
            "focus": "Integrating monetary transactions, loan disbursements, and audit trails.",
            "objective": "Implement financial record-keeping, automated calculation engines, and system-wide action tracking for regulatory compliance.",
            "deliverables": [
                "Financial Ledger and Loan Management module supporting principal, interest, and payment tracking.",
                "Accounts Payable / Accounts Receivable tracking synced with PO status updates.",
                "Centralized Audit Trail logging user actions, IP addresses, timestamped state changes, and transactional diffs.",
                "Dynamic Seeding Command (python manage.py seed_demo) for generating realistic financial and operational test scenarios."
            ],
            "criteria": [
                "Financial calculation logic verified for edge cases (zero interest, early repayment, overpayments).",
                "Audit logs record all CREATE, UPDATE, and DELETE events across critical models.",
                "Data seeding tool cleanly populates realistic demo environments without database integrity violations."
            ]
        },
        {
            "num": "Milestone 4",
            "title": "User Interface, Accessibility & Theme System Refinement",
            "focus": "Delivering a high-impact, modern visual aesthetic with seamless Light/Dark mode toggling.",
            "objective": "Polish the web application user interface to deliver an enterprise-grade experience adhering to accessibility standards.",
            "deliverables": [
                "Modern CSS Tokens & Palette (htc-theme.css) supporting dynamic variables for background, text, borders, and active highlights.",
                "Seamless Theme Switcher maintaining session/local storage state across page reloads.",
                "High-contrast UI components for tables, modals, badges, navigation sidebars, and analytical dashboards.",
                "Asset cache-busting implementation (?v=11 asset versioning) to guarantee client-side CSS synchronization."
            ],
            "criteria": [
                "Zero white-washing or unreadable text when toggled into Dark Mode.",
                "Active row selections, form controls, and badges display WCAG-compliant contrast ratios in both themes.",
                "Dynamic layouts adjust seamlessly across desktop and tablet screen dimensions."
            ]
        },
        {
            "num": "Milestone 5",
            "title": "Automated E2E Testing, Stress Testing (KaneAI) & Security Hardening",
            "focus": "Verifying stability under high load, edge-case vulnerability auditing, and automated regression testing.",
            "objective": "Leverage KaneAI (LambdaTest) cloud suite to subject core business functions to high-stress and multi-user concurrency testing.",
            "deliverables": [
                "Integrated KaneAI automated testing suites for core user flows (Purchase Order Creation, Dark Mode persistence, Loan Disbursements).",
                "Concurrency & Stress Testing scripts validating race condition protection during simultaneous stock updates.",
                "Security Audit matrix covering XSS, CSRF protection, SQL injection resilience, and session fixation defense."
            ],
            "criteria": [
                "All automated KaneAI end-to-end test scenarios achieve green pass status.",
                "Core transaction pipelines maintain data integrity under simulated multi-user concurrent traffic.",
                "Zero critical or high security vulnerabilities identified during code audit."
            ]
        },
        {
            "num": "Milestone 6",
            "title": "Infrastructure Containerization, AWS EC2 Cloud Deployment & Final Handoff",
            "focus": "Containerizing application stacks, orchestrating cloud deployment, and securing production endpoints.",
            "objective": "Deploy the HTC Core application to AWS Cloud infrastructure with production-ready reverse proxies, SSL/TLS, and automated container lifecycle management.",
            "deliverables": [
                "Production Docker configuration (Docker-compose.yml, Dockerfile) encapsulating Django, Gunicorn, PostgreSQL, and Nginx.",
                "AWS EC2 Cloud Instance provisioning with customized security groups and port forwarding (80/443).",
                "Automated migration, static file collection, and superuser initialization scripts.",
                "Production operational runbook detailing backup procedures, log rotation, and server reboot recovery."
            ],
            "criteria": [
                "Application successfully hosted and accessible on live AWS EC2 IP / Domain with sub-second response times.",
                "Production database isolated and persistent across container restarts.",
                "Final system walkthrough validated and signed off by stakeholders and academic project advisor."
            ]
        }
    ]

    for m in milestones_data:
        m_head = doc.add_heading(level=2)
        run_mh = m_head.add_run(f"📌 {m['num']}: {m['title']}")
        run_mh.font.color.rgb = SECONDARY_COLOR
        run_mh.font.size = Pt(12)
        run_mh.font.bold = True

        p_focus = doc.add_paragraph()
        p_focus.paragraph_format.space_after = Pt(4)
        r_f_lbl = p_focus.add_run("Focus: ")
        r_f_lbl.bold = True
        r_f_lbl.font.color.rgb = PRIMARY_COLOR
        p_focus.add_run(m['focus']).italic = True

        p_obj = doc.add_paragraph()
        p_obj.paragraph_format.space_after = Pt(6)
        r_o_lbl = p_obj.add_run("Objective: ")
        r_o_lbl.bold = True
        p_obj.add_run(m['objective'])

        # Key Deliverables List
        p_del_title = doc.add_paragraph()
        p_del_title.paragraph_format.space_after = Pt(2)
        p_del_title.add_run("Key Deliverables:").bold = True

        for item in m['deliverables']:
            p_bullet = doc.add_paragraph(style='List Bullet')
            p_bullet.paragraph_format.space_after = Pt(2)
            p_bullet.add_run(item)

        # Acceptance Criteria
        p_crit_title = doc.add_paragraph()
        p_crit_title.paragraph_format.space_before = Pt(4)
        p_crit_title.paragraph_format.space_after = Pt(2)
        p_crit_title.add_run("Acceptance & Sign-off Criteria:").bold = True

        for item in m['criteria']:
            p_crit = doc.add_paragraph(style='List Bullet')
            p_crit.paragraph_format.space_after = Pt(2)
            run_chk = p_crit.add_run("[✓] ")
            run_chk.font.color.rgb = RGBColor(30, 140, 50)
            run_chk.bold = True
            p_crit.add_run(item)

        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # --- Section: Summary Matrix Table ---
    h3 = doc.add_heading(level=1)
    run_h3 = h3.add_run("3. Milestone Summary & Verification Matrix")
    run_h3.font.color.rgb = PRIMARY_COLOR
    run_h3.font.size = Pt(14)
    run_h3.font.bold = True

    table = doc.add_table(rows=1, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Milestone", "Deliverable Phase", "Primary Target", "Verification Method", "Status"]
    col_widths = [Inches(1.0), Inches(1.4), Inches(1.5), Inches(1.6), Inches(1.0)]

    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "0F2240")
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    matrix_rows = [
        ("M1", "Base Architecture & Security", "System Core, Database Schema, RBAC", "Automated Tests & Route Audits", "COMPLETED"),
        ("M2", "Operations & Logistics", "Procurement, POs, Inventory Tracking", "End-to-End Workflow Verification", "COMPLETED"),
        ("M3", "Finance & Audit Engine", "Loans, Ledger, Activity Auditing", "Calculation Audits & Seed Script", "COMPLETED"),
        ("M4", "UI/UX & Dark Mode Theme", "Modern Design System, Cache Busting", "UI/UX Contrast & Browser Audits", "COMPLETED"),
        ("M5", "KaneAI E2E & Stress Test", "Automated Suites, Concurrency", "KaneAI Cloud Runner Execution", "COMPLETED"),
        ("M6", "AWS EC2 Cloud Deployment", "Docker, Nginx, Production Handoff", "Live EC2 Smoke Test & Review", "FINAL SIGN-OFF"),
    ]

    for r_idx, row_data in enumerate(matrix_rows):
        row_cells = table.add_row().cells
        bg_color = LIGHT_BG if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = val
            row_cells[c_idx].width = col_widths[c_idx]
            set_cell_background(row_cells[c_idx], bg_color)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 4] else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                if c_idx == 4:
                    run.font.bold = True
                    if val == "COMPLETED":
                        run.font.color.rgb = RGBColor(20, 120, 40)
                    else:
                        run.font.color.rgb = RGBColor(30, 90, 180)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)

    # --- Section: Stakeholder Sign-Off Block ---
    h4 = doc.add_heading(level=1)
    run_h4 = h4.add_run("4. Stakeholder & Panelist Approval Sign-Off")
    run_h4.font.color.rgb = PRIMARY_COLOR
    run_h4.font.size = Pt(14)
    run_h4.font.bold = True

    p_sign_intro = doc.add_paragraph()
    p_sign_intro.paragraph_format.space_after = Pt(12)
    p_sign_intro.add_run(
        "By signing below, the undersigned stakeholders, project advisers, and evaluation panelists confirm "
        "that the milestone deliverables detailed above have been reviewed, demonstrated, and accepted in full "
        "according to specified project standards."
    )

    sign_table = doc.add_table(rows=1, cols=4)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False

    sign_headers = ["Stakeholder Role", "Name & Title", "Signature", "Date"]
    sign_col_widths = [Inches(1.6), Inches(2.0), Inches(1.6), Inches(1.3)]

    s_hdr_cells = sign_table.rows[0].cells
    for i, title in enumerate(sign_headers):
        s_hdr_cells[i].text = title
        set_cell_background(s_hdr_cells[i], "1E5A96")
        s_hdr_cells[i].width = sign_col_widths[i]
        p = s_hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)

    sign_roles = [
        "Lead Capstone Adviser",
        "Panelist Evaluator 1",
        "Panelist Evaluator 2",
        "Lead Developer / Author"
    ]

    for role in sign_roles:
        r_cells = sign_table.add_row().cells
        r_cells[0].text = role
        set_cell_background(r_cells[0], "F4F6F9")
        r_cells[0].paragraphs[0].runs[0].font.bold = True
        r_cells[0].paragraphs[0].runs[0].font.size = Pt(9.5)

        for c_idx in range(1, 4):
            r_cells[c_idx].text = ""
            r_cells[c_idx].width = sign_col_widths[c_idx]

    doc.save(output_path)
    print(f"Successfully generated DOCX at {output_path}")

if __name__ == "__main__":
    create_milestones_doc(r"C:\Users\Gigabyte\CAPSTONE-2\HTC_Core_Project_Milestones.docx")
